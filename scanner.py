from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path
from uuid import uuid4

import fitz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


PALETTE = {
	"blue": "#425B9A",
	"sky": "#76C0EC",
	"cream": "#FFF6DC",
	"pink": "#FF95A5",
}

FORMAT_EXTENSIONS = {
	"TIFF": {".tif", ".tiff"},
	"PNG": {".png"},
	"JPG": {".jpg", ".jpeg"},
	"PSD": {".psd"},
	"PDF": {".pdf"},
}
PACKSHOT_EXTENSIONS = set().union(*FORMAT_EXTENSIONS.values()) - {".pdf"}
PACKSHOT_EXTENSIONS.update(
	{
		".avif",
		".arw",
		".bmp",
		".cr2",
		".cr3",
		".dng",
		".gif",
		".heic",
		".heif",
		".ico",
		".jfif",
		".jif",
		".jpe",
		".jp2",
		".j2k",
		".nef",
		".orf",
		".pef",
		".rw2",
		".srw",
		".svg",
		".tga",
		".webp",
	}
)

PDF_FIELD_ORDER = [
	"IDH no.",
	"EAN",
	"Country",
	"Material no.",
	"Article Short Code",
	"SEAL ID",
	"Project No.",
	"File Name",
]

VIEW_PATTERN = re.compile(
	r"(?:^|_)(left_angle|right_angle|front|back|side|left|right|top|bottom|notino|dm|ing)$",
	re.IGNORECASE,
)
NUMERIC_TOKEN_PATTERN = re.compile(r"(?<![A-Za-z0-9])\d+(?![A-Za-z0-9])")
CAPACITY_PATTERN = re.compile(
	r"(?<![A-Za-z0-9])"
	r"(\d+(?:[.,]\d+)?(?:[x×]\d+(?:[.,]\d+)?)?)"
	r"(?:[_\s-]*)"
	r"(fl\.?[_\s-]*oz|millilit(?:er|re)s?|kilograms?|"
	r"centilit(?:er|re)s?|decilit(?:er|re)s?|milligrams?|micrograms?|"
	r"grams?|lit(?:er|re)s?|ounces?|gallons?|quarts?|pints?|"
	r"lbs?|oz|kg|mg|mcg|μg|µg|ml|cl|dl|ul|cc|cm3|gal|qt|pt|k|m|l|g)"
	r"(?![A-Za-z0-9])",
	re.IGNORECASE,
)


def extract_codes_and_view(image_name: str) -> tuple[str, str, str]:
	"""Extract the first EAN and IDH numeric tokens and recognized view suffix."""
	stem = Path(image_name).stem
	numeric_tokens = NUMERIC_TOKEN_PATTERN.findall(stem)
	ean = next((token for token in numeric_tokens if len(token) >= 8), "")
	idh = next((token for token in numeric_tokens if 5 <= len(token) <= 7), "")
	view_match = VIEW_PATTERN.search(stem)
	return (
		ean,
		idh,
		view_match.group(1).lower() if view_match else "",
	)


def extract_capacity(image_name: str) -> str:
	"""Extract a number with a recognized volume or mass unit from an image name."""
	match = CAPACITY_PATTERN.search(Path(image_name).name)
	if not match:
		return ""
	return f"{match.group(1)}{match.group(2).replace('_', '').replace(' ', '')}"


def read_naming_columns(naming_file: Path) -> list[tuple[str, int]]:
	"""Read header names and zero-based column indices from the active worksheet."""
	from openpyxl import load_workbook

	workbook = load_workbook(naming_file, read_only=True, data_only=True)
	try:
		worksheet = workbook.active
		headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
		return [
			(str(value).strip(), index)
			for index, value in enumerate(headers)
			if value is not None and str(value).strip()
		]
	finally:
		workbook.close()


def _normalize_packshot_name(name: str) -> str:
	cleaned = name.strip()
	if Path(cleaned).suffix.casefold() in PACKSHOT_EXTENSIONS:
		return Path(cleaned).stem
	return cleaned


def _rename_packshot(source: Path, target: Path) -> None:
	if source.name.casefold() != target.name.casefold() or source.name == target.name:
		source.rename(target)
		return

	temporary_path = source.with_name(f".packshot_rename_{uuid4().hex}{source.suffix}")
	source.rename(temporary_path)
	try:
		temporary_path.rename(target)
	except Exception:
		temporary_path.rename(source)
		raise


def rename_packshots(
	image_folder: Path,
	naming_file: Path,
	original_column: int,
	new_column: int,
	create_backup: bool,
	export_report: bool,
) -> tuple[Path | None, int, int]:
	"""Rename matching image stems while preserving extensions and report mapping results."""
	from openpyxl import load_workbook

	if not image_folder.is_dir():
		raise ValueError(f"Image folder does not exist: {image_folder}")
	if not naming_file.is_file():
		raise ValueError(f"Naming workbook does not exist: {naming_file}")
	if original_column < 0 or new_column < 0 or original_column == new_column:
		raise ValueError("Choose two different, valid name columns.")

	workbook = load_workbook(naming_file, read_only=True, data_only=True)
	try:
		worksheet = workbook.active
		rows = worksheet.iter_rows(values_only=True)
		headers = next(rows, ())
		if original_column >= len(headers) or new_column >= len(headers):
			raise ValueError("The selected name columns are not available in the naming workbook.")

		mappings: list[tuple[str, str, str, str]] = []
		seen_originals: dict[str, tuple[str, int]] = {}
		for row_number, row in enumerate(rows, start=2):
			original_value = row[original_column] if original_column < len(row) else None
			new_value = row[new_column] if new_column < len(row) else None
			if original_value is None and new_value is None:
				continue
			original_name = str(original_value).strip() if original_value is not None else ""
			new_name = str(new_value).strip() if new_value is not None else ""
			if not original_name or not new_name:
				raise ValueError(
					f"Row {row_number} in the naming workbook must contain both selected names."
				)
			normalized_original = _normalize_packshot_name(original_name)
			normalized_new = _normalize_packshot_name(new_name)
			if not normalized_original or not normalized_new:
				raise ValueError(f"Row {row_number} contains an empty packshot name.")
			key = normalized_original.casefold()
			previous_mapping = seen_originals.get(key)
			if previous_mapping is not None:
				if previous_mapping[0].casefold() != normalized_new.casefold():
					raise ValueError(
						f"Original name '{original_name}' has conflicting replacements "
						f"in rows {previous_mapping[1]} and {row_number}."
					)
				continue
			seen_originals[key] = (normalized_new, row_number)
			mappings.append(
				(original_name, new_name, normalized_original, normalized_new)
			)
	finally:
		workbook.close()

	image_paths = sorted(
		(
			path
			for path in image_folder.rglob("*")
			if path.is_file()
			and path.suffix.casefold() in PACKSHOT_EXTENSIONS
			and (not path.relative_to(image_folder).parts
				or path.relative_to(image_folder).parts[0].casefold() != "backup")
		),
		key=lambda path: str(path).casefold(),
	)
	images_by_stem: dict[str, list[Path]] = {}
	for image_path in image_paths:
		images_by_stem.setdefault(image_path.stem.casefold(), []).append(image_path)

	max_folder_depth = max(
		(
			len(path.relative_to(image_folder).parts) - 1
			for path in image_paths
		),
		default=0,
	)
	folder_headers = [f"folder {level}" for level in range(1, max_folder_depth + 2)]
	report_headers = [
		*folder_headers,
		"idh",
		"ean",
		"capacity",
		"original_name",
		"new_name",
		"is_original_name_replaced?",
		"is_new_name_used?",
	]
	report_rows: list[list[str]] = []
	renamed_count = 0
	skipped_count = 0
	backup_root = image_folder / "backup"

	for original_name, new_name, normalized_original, normalized_new in mappings:
		matches = [
			path
			for path in images_by_stem.get(normalized_original.casefold(), [])
			if path.stem.casefold() != normalized_new.casefold()
		]
		already_named_images = images_by_stem.get(normalized_new.casefold(), [])
		for image_path in already_named_images:
			relative_path = image_path.relative_to(image_folder)
			folders = [image_folder.name, *relative_path.parts[:-1]]
			folder_columns = [*folders, *("" for _ in range(len(folder_headers) - len(folders)))]
			ean, idh, _ = extract_codes_and_view(image_path.stem)
			capacity = extract_capacity(image_path.name)
			report_rows.append(
				[
					*folder_columns,
					idh,
					ean,
					capacity,
					original_name,
					new_name,
					"no",
					"yes",
				]
			)

		if not matches:
			if already_named_images:
				continue
			report_rows.append(
				[
					*("" for _ in folder_headers),
					"",
					"",
					"",
					original_name,
					new_name,
					"no",
					"no",
				]
			)
			continue

		for image_path in matches:
			relative_path = image_path.relative_to(image_folder)
			folders = [image_folder.name, *relative_path.parts[:-1]]
			folder_columns = [*folders, *("" for _ in range(len(folder_headers) - len(folders)))]
			ean, idh, _ = extract_codes_and_view(image_path.stem)
			capacity = extract_capacity(image_path.name)
			target_path = image_path.with_name(f"{normalized_new}{image_path.suffix}")
			same_name = image_path.name == target_path.name
			case_only_rename = image_path.name.casefold() == target_path.name.casefold()
			new_name_used = bool(already_named_images)

			if same_name:
				new_name_used = True
			elif not case_only_rename and target_path.exists():
				skipped_count += 1
				new_name_used = True
			elif create_backup:
				backup_path = backup_root / relative_path
				if backup_path.exists():
					skipped_count += 1
				else:
					backup_path.parent.mkdir(parents=True, exist_ok=True)
					shutil.copy2(image_path, backup_path)
					_rename_packshot(image_path, target_path)
					renamed_count += 1
					new_name_used = True
			else:
				_rename_packshot(image_path, target_path)
				renamed_count += 1
				new_name_used = True

			report_rows.append(
				[
					*folder_columns,
					idh,
					ean,
					capacity,
					original_name,
					new_name,
					"yes",
					"yes" if new_name_used else "no",
				]
			)

	if not export_report:
		return None, renamed_count, skipped_count

	timestamp = datetime.now().strftime("%Y_%m_%d_%H_%M_%S")
	report_path = naming_file.parent / f"packshot_rename_report_{timestamp}.xlsx"
	_save_excel_report(report_path, report_headers, report_rows)
	return report_path, renamed_count, skipped_count


def _clean_pdf_value(raw_value: str) -> str:
	cleaned = raw_value.replace("\u00a0", " ")
	cleaned = re.sub(r"\s+", " ", cleaned)
	cleaned = cleaned.strip(" :;,-")
	return cleaned.strip()


def _extract_pdf_field(pdf_text: str, label: str) -> str:
	patterns = {
		"IDH no.": re.compile(r"IDH\s*no\.?\s*(?:[:\-]?\s*)?(\d{5,8})", re.IGNORECASE),
		"EAN": re.compile(
			r"EAN(?:\s*\(please check\))?\s*(?:[:\-]?\s*)?(\d{8,18})",
			re.IGNORECASE,
		),
		"Country": re.compile(
			r"Country\s*(?:[:\-]?\s*)?([A-Z0-9:\(\)\-\s/]+?)(?=(?:\s*(?:Creation\s+date|Product\s+|Status\s+|Date\s+of\s+approval|$)))",
			re.IGNORECASE,
		),
		"Material no.": re.compile(r"Material\s*no\.?\s*(?:[:\-]?\s*)?([A-Z0-9\-]+)", re.IGNORECASE),
		"Article Short Code": re.compile(
			r"Article\s+Short\s+code\s*(?:[:\-]?\s*)?([A-Z0-9\-]+)",
			re.IGNORECASE,
		),
		"SEAL ID": re.compile(r"SEAL\s*ID\s*(?:[:\-]?\s*)?([A-Z0-9\-]+)", re.IGNORECASE),
		"Project No.": re.compile(
			r"Project\s*no\.?\s*(?:\(\s*HINT\s*\))?\s*(?:[:\-]?\s*)?([A-Z0-9\-]+)",
			re.IGNORECASE,
		),
		"File Name": re.compile(
			r"(?:LSD\s+FILE\s+NAME|File\s+Name)\s*(?:[:\-]?\s*)?([A-Za-z0-9_\-\.]+)",
			re.IGNORECASE,
		),
	}
	label_pattern = patterns.get(label)
	if label_pattern is None:
		return ""
	match = label_pattern.search(pdf_text)
	if not match:
		return ""
	value = match.group(1)
	if label == "EAN":
		value = re.sub(r"^0+", "", value)
	value = _clean_pdf_value(value)
	return value


def extract_pdf_data(pdf_path: Path) -> dict[str, str]:
	"""Read key-value data from the PDF info box and return selected fields."""
	text_chunks: list[str] = []
	try:
		with fitz.open(pdf_path) as document:
			for page in document:
				text_chunks.append(page.get_text("text", sort=True))
			pdf_text = " ".join(text_chunks)
	except Exception as error:
		raise RuntimeError(f"Unable to read PDF metadata from '{pdf_path}': {error}") from error

	fields = {
		field: _extract_pdf_field(pdf_text, field)
		for field in PDF_FIELD_ORDER
	}
	if not fields["File Name"]:
		fields["File Name"] = pdf_path.name
	return fields


def scan_images(
	source_folder: Path,
	selected_formats: set[str],
	include_extension: bool,
) -> tuple[list[list[str]], list[str]]:
	"""Return workbook rows and headers for matching images under source_folder."""
	suffixes = {
		suffix
		for image_format in selected_formats
		for suffix in FORMAT_EXTENSIONS[image_format]
	}
	matches: list[tuple[Path, tuple[str, ...]]] = []

	for image_path in source_folder.rglob("*"):
		if image_path.is_file() and image_path.suffix.lower() in suffixes:
			relative_path = image_path.relative_to(source_folder)
			matches.append((image_path, relative_path.parts[:-1]))

	matches.sort(key=lambda match: str(match[0]).casefold())
	deepest_folder = max((len(folders) for _, folders in matches), default=0)
	headers = [f"folder {level}" for level in range(1, deepest_folder + 2)]
	headers.extend(["ean", "idh", "capacity", "image_name", "image_format", "view"])

	rows: list[list[str]] = []
	for image_path, folders in matches:
		folder_columns = [source_folder.name, *folders]
		folder_columns.extend([""] * (deepest_folder + 1 - len(folder_columns)))
		ean, idh, view = extract_codes_and_view(image_path.stem)
		capacity = extract_capacity(image_path.stem)
		image_name = image_path.name if include_extension else image_path.stem
		image_format = next(
			name
			for name, extensions in FORMAT_EXTENSIONS.items()
			if image_path.suffix.lower() in extensions
		)
		rows.append([*folder_columns, ean, idh, capacity, image_name, image_format, view])

	return rows, headers


def scan_pdf_files(
	source_folder: Path,
	selected_fields: set[str],
) -> tuple[list[list[str]], list[str]]:
	"""Return pdf rows and headers for selected metadata fields."""
	pdf_matches: list[Path] = []
	for pdf_path in source_folder.rglob("*"):
		if pdf_path.is_file() and pdf_path.suffix.lower() == ".pdf":
			pdf_matches.append(pdf_path)

	pdf_matches.sort(key=lambda path: str(path).casefold())
	deepest_folder = max(
		(len(pdf_path.relative_to(source_folder).parts[:-1]) for pdf_path in pdf_matches),
		default=0,
	)
	headers = [f"folder {level}" for level in range(1, deepest_folder + 2)]
	headers.extend(field for field in PDF_FIELD_ORDER if field in selected_fields)

	rows: list[list[str]] = []
	for pdf_path in pdf_matches:
		relative_path = pdf_path.relative_to(source_folder)
		folders = relative_path.parts[:-1]
		folder_columns = [source_folder.name, *folders]
		folder_columns.extend([""] * (deepest_folder + 1 - len(folder_columns)))
		fields = extract_pdf_data(pdf_path)
		row: list[str] = [*folder_columns]
		for field in PDF_FIELD_ORDER:
			if field in selected_fields:
				row.append(fields.get(field, ""))
		rows.append(row)

	return rows, headers


def _save_excel_report(
	report_path: Path,
	headers: list[str],
	rows: list[list[str]],
) -> None:
	workbook = Workbook()
	worksheet = workbook.active
	worksheet.title = "report"
	worksheet.append(headers)
	for row in rows:
		worksheet.append(row)

	header_fill = PatternFill("solid", fgColor=PALETTE["blue"].lstrip("#"))
	for cell in worksheet[1]:
		cell.fill = header_fill
		cell.font = Font(color="FFFFFF", bold=True)
	worksheet.freeze_panes = "A2"
	worksheet.auto_filter.ref = worksheet.dimensions
	for column in worksheet.columns:
		column_letter = column[0].column_letter
		max_width = max(len(str(cell.value or "")) for cell in column)
		worksheet.column_dimensions[column_letter].width = min(max(max_width + 2, 14), 42)
	workbook.save(report_path)


def write_report(
	source_folder: Path,
	output_folder: Path,
	selected_formats: set[str],
	include_extension: bool,
) -> tuple[Path, int]:
	rows, headers = scan_images(source_folder, selected_formats, include_extension)
	timestamp = datetime.now().strftime("%Y_%m_%d_%H_%M_%S")
	report_path = output_folder / f"scanned_packshots_{timestamp}.xlsx"

	workbook = Workbook()
	worksheet = workbook.active
	worksheet.title = "report"
	worksheet.append(headers)
	for row in rows:
		worksheet.append(row)

	header_fill = PatternFill("solid", fgColor=PALETTE["blue"].lstrip("#"))
	for cell in worksheet[1]:
		cell.fill = header_fill
		cell.font = Font(color="FFFFFF", bold=True)
	worksheet.freeze_panes = "A2"
	worksheet.auto_filter.ref = worksheet.dimensions
	for column in worksheet.columns:
		column_letter = column[0].column_letter
		max_width = max(len(str(cell.value or "")) for cell in column)
		worksheet.column_dimensions[column_letter].width = min(max(max_width + 2, 14), 42)

	workbook.save(report_path)
	return report_path, len(rows)


def write_pdf_report(
	source_folder: Path,
	output_folder: Path,
	selected_fields: set[str],
) -> tuple[Path, int]:
	rows, headers = scan_pdf_files(source_folder, selected_fields)
	timestamp = datetime.now().strftime("%Y_%m_%d_%H_%M")
	report_path = output_folder / f"pdf_data_{timestamp}.xlsx"

	workbook = Workbook()
	worksheet = workbook.active
	worksheet.title = "report"
	worksheet.append(headers)
	for row in rows:
		worksheet.append(row)

	header_fill = PatternFill("solid", fgColor=PALETTE["blue"].lstrip("#"))
	for cell in worksheet[1]:
		cell.fill = header_fill
		cell.font = Font(color="FFFFFF", bold=True)
	worksheet.freeze_panes = "A2"
	worksheet.auto_filter.ref = worksheet.dimensions
	for column in worksheet.columns:
		column_letter = column[0].column_letter
		max_width = max(len(str(cell.value or "")) for cell in column)
		worksheet.column_dimensions[column_letter].width = min(max(max_width + 2, 14), 42)

	workbook.save(report_path)
	return report_path, len(rows)
